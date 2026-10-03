"""Generate sheet bergaya 'PENDIDIKAN' (surat pernyataan) dari template DUPAK
asli -- meniru struktur & teks aslinya (dikelompokkan per semester, sub-total
berjenjang, formula antar baris), bukan isi ulang seluruh form resmi (banyak
bidang yang belum di-pipeline, misal pendidikan formal/ijazah, disertasi,
thesis) -- hanya 3 kategori yang sudah divalidasi dari data nyata: Mengajar,
Menguji, Bimbingan Publikasi Karya Ilmiah.

ASUMSI PEMETAAN yang perlu direview manual sebelum dipakai resmi:
  - "Menguji" (publikasi karya ilmiah) dipetakan ke sub-baris "c. Skripsi"
    pada "Bertugas sebagai penguji" -- karena publikasi karya ilmiah jalur
    non-skripsi adalah pengganti skripsi, bukan kategori terpisah di form resmi.
  - "Bimbingan" (publikasi karya ilmiah) dipetakan ke sub-baris "d. Laporan
    akhir studi" pada "Membimbing ... Pembimbing Utama" -- dikonfirmasi oleh
    dosen pemilik data sebagai setara skripsi, pembimbing tunggal/utama.
  - Semester untuk Menguji & Bimbingan DIHITUNG dari tanggal dokumen (lihat
    tanggal_indo.py), karena sumber datanya cuma punya tanggal, bukan label
    semester eksplisit seperti SK Mengajar. Kalau tanggalnya tidak kebaca,
    aktivitas itu dikelompokkan ke "Tanggal tidak diketahui" di akhir.
"""

from __future__ import annotations

import json
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from agentdupak.tanggal_indo import parse_tanggal, rentang_bulan, tentukan_semester

FONT_NORMAL = Font(name="Arial", size=10)
FONT_BOLD = Font(name="Arial", size=10, bold=True)
FONT_JUDUL = Font(name="Arial", size=12, bold=True)
FONT_HEADER = Font(name="Arial", size=10, bold=True, color="FFFFFF")
FILL_HEADER = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")

# posisi kolom data (1-indexed): A No, B Uraian level-1, C level-2, D level-3, E level-4,
# H Tanggal, I Satuan Hasil, J Jumlah Volume, K Angka Kredit (satuan), L Jumlah AK, M Keterangan
COL_NO, COL_B, COL_C, COL_D, COL_E = 1, 2, 3, 4, 5
COL_TANGGAL, COL_SATUAN, COL_VOLUME, COL_AK_SATUAN, COL_JUMLAH_AK, COL_KET = 8, 9, 10, 11, 12, 13

TANPA_TANGGAL = "Tanggal tidak diketahui"


def _semester_group_key(tanggal_teks: str | None) -> tuple:
    d = parse_tanggal(tanggal_teks)
    if d is None:
        return (9999, 9, TANPA_TANGGAL, "")
    label, tahun_ajaran = tentukan_semester(d)
    tahun_awal = int(tahun_ajaran.split("/")[0])
    urutan = 0 if label == "Ganjil" else 1
    return (tahun_awal, urutan, label, tahun_ajaran)


def _kelompokkan_per_semester(aktivitas: list[dict]) -> list[tuple[str, str, list[dict]]]:
    """Kembalikan [(label, tahun_ajaran, [aktivitas...]), ...] terurut kronologis."""
    kelompok: dict[tuple, list[dict]] = defaultdict(list)
    for a in aktivitas:
        key = _semester_group_key(a["atribut"].get("tanggal"))
        kelompok[key].append(a)

    hasil = []
    for key in sorted(kelompok.keys()):
        _, _, label, tahun_ajaran = key
        hasil.append((label, tahun_ajaran, kelompok[key]))
    return hasil


def _set_lebar_kolom(ws: Worksheet) -> None:
    lebar = {1: 5, 2: 6, 3: 6, 4: 6, 5: 38, 8: 16, 9: 10, 10: 10, 11: 10, 12: 10, 13: 20}
    for kolom, w in lebar.items():
        ws.column_dimensions[get_column_letter(kolom)].width = w


def _tulis_header_kolom(ws: Worksheet, baris: int) -> None:
    posisi = {1: "No.", 2: "Uraian Kegiatan", 8: "Tanggal", 9: "Satuan Hasil",
              10: "Jumlah\nVolume\nKegiatan", 11: "Angka Kredit", 12: "Jumlah\nAngka\nKredit",
              13: "Keterangan/\nBukti Fisik"}
    for col, teks in posisi.items():
        c = ws.cell(row=baris, column=col, value=teks)
        c.font = FONT_HEADER
        c.fill = FILL_HEADER
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")


def generate_workbook(dosen_nama: str, aktivitas_rows) -> Workbook:
    """aktivitas_rows: hasil db.get_aktivitas() -- list sqlite3.Row."""
    mengajar = []
    menguji = []
    bimbingan = []
    lainnya = []
    for row in aktivitas_rows:
        atribut = json.loads(row["atribut_json"])
        item = {"atribut": atribut, "ak": row["ak"], "ak_perlu_review": row["ak_perlu_review"], "sumber_file": row["sumber_file"]}
        if row["kategori"] == "pendidikan.mengajar":
            mengajar.append(item)
        elif row["kategori"] == "pendidikan.menguji":
            menguji.append(item)
        elif row["kategori"] == "pendidikan.bimbingan_publikasi_ilmiah":
            bimbingan.append(item)
        else:
            lainnya.append((row["kategori"], item))

    wb = Workbook()
    ws = wb.active
    ws.title = "PENDIDIKAN"
    _set_lebar_kolom(ws)

    ws["A1"] = "SURAT PERNYATAAN"
    ws["A1"].font = FONT_JUDUL
    ws["A2"] = "MELAKSANAKAN PENDIDIKAN"
    ws["A2"].font = FONT_JUDUL
    ws["A4"] = f"Dihasilkan otomatis oleh agentdupak untuk: {dosen_nama}"
    ws["A4"].font = FONT_NORMAL
    ws["A5"] = (
        "Field bio dosen (NIDN, pangkat, unit kerja) belum diisi otomatis -- "
        "lengkapi manual sebelum dipakai resmi."
    )
    ws["A5"].font = Font(name="Arial", size=9, italic=True, color="888888")

    baris = 7
    _tulis_header_kolom(ws, baris)
    baris += 2

    # ================= A. Melaksanakan perkuliahan (mengajar) =================
    baris_header_mengajar = baris
    ws.cell(row=baris, column=COL_B, value="A").font = FONT_BOLD
    ws.cell(row=baris, column=COL_C, value=(
        "Melaksanakan perkulihan/ tutorial dan membimbing, menguji serta menyelenggarakan "
        "pendidikan di laboratorium, praktek keguruan bengkel/ studio/kebun pada "
        "Fakultas/Sekolah Tinggi/Akademi/ Politeknik sendiri, pada fakultas lain dalam "
        "lingkungan Universitas/Institut sendiri, maupun di luar perguruan tinggi sendiri "
        "secara melembaga paling banyak 12 sks per semester"
    )).font = FONT_NORMAL
    ws.cell(row=baris, column=COL_C).alignment = Alignment(wrap_text=True, vertical="top")
    baris += 1

    kelompok_mengajar = []
    for atribut_key in sorted({(a["atribut"].get("tahun_ajaran", ""), a["atribut"].get("semester", "")) for a in mengajar}):
        tahun_ajaran, label = atribut_key
        anggota = [a for a in mengajar if a["atribut"].get("tahun_ajaran") == tahun_ajaran and a["atribut"].get("semester") == label]
        if anggota:
            kelompok_mengajar.append((label, tahun_ajaran, anggota))
    # urutkan kronologis (ganjil sebelum genap dalam 1 tahun ajaran yang sama tidak selalu
    # benar urutan kalender -- tapi cukup untuk tampilan, bukan penentu nilai AK)
    kelompok_mengajar.sort(key=lambda t: (t[1], 0 if t[0] == "Ganjil" else 1))

    baris_subtotal_mengajar = []
    nomor_semester = 1
    for label, tahun_ajaran, anggota in kelompok_mengajar:
        rentang = rentang_bulan(label, tahun_ajaran) if label in ("Genap", "Ganjil") else ""
        ws.cell(row=baris, column=COL_C, value=(
            f"{nomor_semester}. Semester {label.lower()} {tahun_ajaran} {rentang}; maksimum 12 SKS per semester"
        )).font = FONT_NORMAL
        nomor_semester += 1
        baris += 1

        baris_mulai_kelas = baris
        for i, a in enumerate(anggota, start=1):
            atr = a["atribut"]
            ws.cell(row=baris, column=COL_C, value=i).font = FONT_NORMAL
            ws.cell(row=baris, column=COL_D, value=f"{atr.get('mata_kuliah', '?')} ({atr.get('kelas', '?')})").font = FONT_NORMAL
            ws.cell(row=baris, column=COL_TANGGAL, value=f"Semester {label} {tahun_ajaran}").font = FONT_NORMAL
            ws.cell(row=baris, column=COL_SATUAN, value="SKS").font = FONT_NORMAL
            ws.cell(row=baris, column=COL_VOLUME, value=atr.get("sks")).font = FONT_NORMAL
            ws.cell(row=baris, column=COL_AK_SATUAN, value=1).font = FONT_NORMAL
            ws.cell(row=baris, column=COL_JUMLAH_AK, value=f"=J{baris}*K{baris}").font = FONT_NORMAL
            ws.cell(row=baris, column=COL_KET, value=a["sumber_file"]).font = FONT_NORMAL
            baris += 1
        baris_akhir_kelas = baris - 1

        ws.cell(row=baris, column=COL_D, value="Sub total per semester").font = FONT_BOLD
        ws.cell(row=baris, column=COL_VOLUME, value=f"=SUM(J{baris_mulai_kelas}:J{baris_akhir_kelas})").font = FONT_BOLD
        ws.cell(row=baris, column=COL_JUMLAH_AK, value=f"=SUM(L{baris_mulai_kelas}:L{baris_akhir_kelas})").font = FONT_BOLD
        baris_subtotal_mengajar.append(baris)
        baris += 1

    if baris_subtotal_mengajar:
        formula_total_mengajar = "=" + "+".join(f"L{b}" for b in baris_subtotal_mengajar)
    else:
        formula_total_mengajar = 0
    ws.cell(row=baris_header_mengajar, column=COL_JUMLAH_AK, value=formula_total_mengajar).font = FONT_BOLD
    baris += 1

    # ================= D. Membimbing (publikasi karya ilmiah -> Laporan akhir) =================
    baris_header_bimbingan = baris
    ws.cell(row=baris, column=COL_B, value="D").font = FONT_BOLD
    ws.cell(row=baris, column=COL_C, value=(
        "Membimbing dan ikut membimbing dalam menghasilkan disertasi, thesis, skripsi dan "
        "laporan akhir studi yang sesuai bidang penugasannya (maksimum 32 kum per semester) "
        "-- di sini diisi dari Pembimbing Publikasi Karya Ilmiah (jalur non-skripsi), "
        "dipetakan ke sub-baris 'd. Laporan akhir studi', Pembimbing Utama"
    )).font = FONT_NORMAL
    ws.cell(row=baris, column=COL_C).alignment = Alignment(wrap_text=True, vertical="top")
    baris += 1

    baris_subtotal_bimbingan = []
    for label, tahun_ajaran, anggota in _kelompokkan_per_semester(bimbingan):
        label_tampil = f"Semester {label.lower()} {tahun_ajaran}" if label != TANPA_TANGGAL else TANPA_TANGGAL
        ws.cell(row=baris, column=COL_C, value=label_tampil).font = FONT_NORMAL
        baris += 1
        ws.cell(row=baris, column=COL_D, value="Pembimbing Utama per orang (setiap mahasiswa)").font = FONT_NORMAL
        baris += 1
        ws.cell(row=baris, column=COL_E, value="d. Laporan akhir studi (maksimum 10 lulusan per semester)").font = FONT_NORMAL
        baris += 1

        baris_mulai = baris
        for a in anggota:
            atr = a["atribut"]
            nim = atr.get("nim") or "-"
            ws.cell(row=baris, column=COL_E, value=f"{atr.get('nama_mahasiswa', '?')} ({nim})").font = FONT_NORMAL
            ws.cell(row=baris, column=COL_TANGGAL, value=atr.get("tanggal") or "-").font = FONT_NORMAL
            ws.cell(row=baris, column=COL_SATUAN, value="Mahasiswa").font = FONT_NORMAL
            ws.cell(row=baris, column=COL_VOLUME, value=1).font = FONT_NORMAL
            ws.cell(row=baris, column=COL_AK_SATUAN, value=a["ak"]).font = FONT_NORMAL
            ws.cell(row=baris, column=COL_JUMLAH_AK, value=f"=J{baris}*K{baris}").font = FONT_NORMAL
            ws.cell(row=baris, column=COL_KET, value=a["sumber_file"]).font = FONT_NORMAL
            baris += 1
        baris_akhir = baris - 1

        ws.cell(row=baris, column=COL_E, value="Sub total pembimbing utama").font = FONT_BOLD
        ws.cell(row=baris, column=COL_JUMLAH_AK, value=f"=SUM(L{baris_mulai}:L{baris_akhir})").font = FONT_BOLD
        baris_subtotal_bimbingan.append(baris)
        baris += 1

    formula_total_bimbingan = "=" + "+".join(f"L{b}" for b in baris_subtotal_bimbingan) if baris_subtotal_bimbingan else 0
    ws.cell(row=baris_header_bimbingan, column=COL_JUMLAH_AK, value=formula_total_bimbingan).font = FONT_BOLD
    baris += 1

    # ================= E. Bertugas sebagai penguji (-> skripsi/setara) =================
    baris_header_menguji = baris
    ws.cell(row=baris, column=COL_B, value="E").font = FONT_BOLD
    ws.cell(row=baris, column=COL_C, value=(
        "Bertugas sebagai penguji pada ujian akhir/Profesi (maksimum 8 kum per semester) "
        "-- di sini diisi dari Penguji Publikasi Karya Ilmiah (jalur non-skripsi), "
        "dipetakan ke sub-baris 'c. Skripsi'"
    )).font = FONT_NORMAL
    ws.cell(row=baris, column=COL_C).alignment = Alignment(wrap_text=True, vertical="top")
    baris += 1

    baris_subtotal_menguji_per_semester = []
    for label, tahun_ajaran, anggota in _kelompokkan_per_semester(menguji):
        label_tampil = f"Semester {label.lower()} {tahun_ajaran}" if label != TANPA_TANGGAL else TANPA_TANGGAL
        ws.cell(row=baris, column=COL_C, value=label_tampil).font = FONT_NORMAL
        baris += 1

        baris_subtotal_peran = []
        for peran, label_peran in (("ketua", "Ketua penguji (maksimum 4 lulusan per semester)"),
                                    ("anggota", "Anggota penguji (maksimum 8 lulusan per semester)")):
            anggota_peran = [a for a in anggota if a["atribut"].get("peran") == peran]
            if not anggota_peran:
                continue
            ws.cell(row=baris, column=COL_D, value=label_peran).font = FONT_NORMAL
            baris += 1
            ws.cell(row=baris, column=COL_E, value="c. Skripsi").font = FONT_NORMAL
            baris += 1

            baris_mulai = baris
            for a in anggota_peran:
                atr = a["atribut"]
                npm = atr.get("npm") or "-"
                ws.cell(row=baris, column=COL_E, value=f"{atr.get('nama_mahasiswa', '?')} ({npm})").font = FONT_NORMAL
                ws.cell(row=baris, column=COL_TANGGAL, value=atr.get("tanggal") or "-").font = FONT_NORMAL
                ws.cell(row=baris, column=COL_SATUAN, value="Mahasiswa").font = FONT_NORMAL
                ws.cell(row=baris, column=COL_VOLUME, value=1).font = FONT_NORMAL
                ws.cell(row=baris, column=COL_AK_SATUAN, value=a["ak"]).font = FONT_NORMAL
                ws.cell(row=baris, column=COL_JUMLAH_AK, value=f"=J{baris}*K{baris}").font = FONT_NORMAL
                ws.cell(row=baris, column=COL_KET, value=a["sumber_file"]).font = FONT_NORMAL
                baris += 1
            baris_akhir = baris - 1

            ws.cell(row=baris, column=COL_E, value=f"Sub total {peran} penguji pada ujian akhir").font = FONT_BOLD
            ws.cell(row=baris, column=COL_JUMLAH_AK, value=f"=SUM(L{baris_mulai}:L{baris_akhir})").font = FONT_BOLD
            baris_subtotal_peran.append(baris)
            baris += 1

        if baris_subtotal_peran:
            ws.cell(row=baris, column=COL_D, value="Sub total per semester").font = FONT_BOLD
            ws.cell(row=baris, column=COL_JUMLAH_AK, value="=" + "+".join(f"L{b}" for b in baris_subtotal_peran)).font = FONT_BOLD
            baris_subtotal_menguji_per_semester.append(baris)
            baris += 1

    formula_total_menguji = "=" + "+".join(f"L{b}" for b in baris_subtotal_menguji_per_semester) if baris_subtotal_menguji_per_semester else 0
    ws.cell(row=baris_header_menguji, column=COL_JUMLAH_AK, value=formula_total_menguji).font = FONT_BOLD
    baris += 2

    # ================= Total keseluruhan =================
    ws.cell(row=baris, column=COL_B, value="JUMLAH (kategori yang sudah di-digitalisasi)").font = FONT_BOLD
    ws.cell(row=baris, column=COL_JUMLAH_AK, value=f"=L{baris_header_mengajar}+L{baris_header_bimbingan}+L{baris_header_menguji}").font = FONT_BOLD

    if lainnya:
        baris += 2
        ws.cell(row=baris, column=COL_B, value=(
            f"Catatan: {len(lainnya)} aktivitas berkategori belum dipetakan ke form ini "
            "(ak_perlu_review) -- lihat sheet 'Belum Terpetakan'."
        )).font = Font(name="Arial", size=9, italic=True, color="C0392B")
        ws_lain = wb.create_sheet("Belum Terpetakan")
        ws_lain.append(["Kategori", "Sumber File", "AK sementara"])
        for kategori, item in lainnya:
            ws_lain.append([kategori, item["sumber_file"], item["ak"]])

    return wb
